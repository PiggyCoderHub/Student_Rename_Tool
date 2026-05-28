# -*- coding: utf-8 -*-
"""
Docx文件读取模块 v2
功能：从docx文件中提取学号、姓名等信息
"""

import re
from pathlib import Path
from typing import Dict, Optional, List, Tuple
from docx import Document


class DocxReader:
    """Word文档读取器 - 增强版"""

    def __init__(self, docx_path: str):
        self.docx_path = Path(docx_path)
        self.text_content = ""
        self.paragraphs = []
        self.tables = []
        self.full_text = ""  # 包括所有文本的完整内容

    def _load_document(self) -> bool:
        """加载文档"""
        try:
            doc = Document(self.docx_path)

            # 提取段落文本
            self.paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            self.text_content = '\n'.join(self.paragraphs)

            # 提取表格内容
            self.tables = []
            for table in doc.tables:
                table_rows = []
                for row in table.rows:
                    row_text = [cell.text.strip() for cell in row.cells]
                    table_rows.append(' | '.join(row_text))
                self.tables.append('\n'.join(table_rows))

            # 合并所有文本用于搜索
            self.full_text = self.text_content + '\n'.join(self.tables)

            return True
        except Exception as e:
            self.full_text = ""
            return False

    def extract_student_id(self) -> Optional[str]:
        """提取学号 - 多种策略"""
        # 策略1：精确匹配表格中的学号
        for table_text in self.tables:
            # 学号: 202301010001 格式
            matches = re.findall(r'(?:学号|Student\s*ID|No\.?|编号)\s*[:：]\s*(\d{8,15})', table_text, re.IGNORECASE)
            if matches:
                return matches[0]

            # 纯12位数字（常见学号格式）
            matches = re.findall(r'\b(\d{12})\b', table_text)
            if matches:
                return matches[0]

            # 以20开头的10位以上数字
            matches = re.findall(r'\b(20\d{8,12})\b', table_text)
            if matches:
                return matches[0]

        # 策略2：从段落中提取
        lines = self.full_text.split('\n')
        for line in lines:
            line = line.strip()
            # 学号: xxx 格式
            match = re.search(r'(?:学号|Student\s*ID|No\.?)\s*[:：]\s*(\d+)', line, re.IGNORECASE)
            if match:
                return match.group(1)

        # 策略3：查找所有数字序列
        all_numbers = re.findall(r'\b(\d{10,14})\b', self.full_text)
        if all_numbers:
            # 返回最长的那个（通常学号位数较多）
            return max(all_numbers, key=len)

        # 策略4：查找文件名中的学号（如果文件名包含）
        filename = self.docx_path.stem
        match = re.search(r'(\d{10,14})', filename)
        if match:
            return match.group(1)

        return None

    def extract_name(self) -> Optional[str]:
        """提取姓名 - 多种策略"""
        # 策略1：从表格中提取
        for table_text in self.tables:
            # 姓名: 张三 格式
            matches = re.findall(r'(?:姓名|Name)\s*[:：]\s*([\u4e00-\u9fa5]{2,4})', table_text, re.IGNORECASE)
            if matches:
                return matches[0]

            # 姓名: Zhang San 格式（英文）
            matches = re.findall(r'(?:姓名|Name)\s*[:：]\s*([A-Za-z]+\s+[A-Za-z]+)', table_text, re.IGNORECASE)
            if matches:
                return matches[0]

        # 策略2：从段落中提取
        lines = self.full_text.split('\n')
        for line in lines:
            line = line.strip()
            # 姓名: xxx 格式
            match = re.search(r'(?:姓名|Name)\s*[:：]\s*([\u4e00-\u9fa5A-Za-z]+)', line, re.IGNORECASE)
            if match:
                return match.group(1)

        # 策略3：从文件名中提取
        filename = self.docx_path.stem
        # 尝试匹配 "姓名" 或 "xxx报告" 模式
        patterns = [
            r'([\u4e00-\u9fa5]{2,4})(?:报告|实验|作业|提交)',  # xxx报告
            r'报告[_ ]?([\u4e00-\u9fa5]{2,4})',  # 报告_xxx
            r'([\u4e00-\u9fa5]{2,4})_',  # xxx_
        ]
        for pattern in patterns:
            match = re.search(pattern, filename)
            if match:
                name = match.group(1)
                # 排除常见关键词
                if name not in ['实验', '报告', '作业', '提交', '文档']:
                    return name

        # 策略4：查找独立的2-4字中文（可能是姓名）
        # 查找 "姓名" 后面的内容
        match = re.search(r'姓名[^\u4e00-\u9fa5]*([\u4e00-\u9fa5]{2,4})', self.full_text)
        if match:
            return match.group(1)

        return None

    def extract_all_info(self) -> Dict[str, Optional[str]]:
        """提取所有信息"""
        self._load_document()

        return {
            'student_id': self.extract_student_id(),
            'name': self.extract_name(),
            'raw_preview': self.full_text[:1000] if self.full_text else None
        }

    def has_valid_info(self) -> bool:
        """检查是否包含有效的学号或姓名"""
        info = self.extract_all_info()
        return bool(info['student_id'] or info['name'])

    def get_all_text(self) -> str:
        """获取文档全部文本"""
        self._load_document()
        return self.full_text


class BatchDocxReader:
    """批量文档读取器"""

    def __init__(self, folder_path: str):
        self.folder_path = Path(folder_path)
        self.docx_files: List[Path] = []
        self.results: List[Dict] = []

    def find_docx_files(self) -> List[Path]:
        """查找文件夹中的docx文件"""
        if not self.folder_path.exists():
            return []

        self.docx_files = list(self.folder_path.glob("*.docx"))
        self.docx_files.extend(self.folder_path.glob("*.doc"))
        return self.docx_files

    def read_all(self) -> List[Dict]:
        """读取所有docx文件"""
        self.results = []

        for docx_file in self.find_docx_files():
            reader = DocxReader(str(docx_file))
            info = reader.extract_all_info()
            info['filename'] = docx_file.name
            info['filepath'] = str(docx_file)
            info['has_info'] = reader.has_valid_info()
            self.results.append(info)

        return self.results

    def get_files_without_info(self) -> List[Dict]:
        """获取无法提取信息的文件"""
        return [r for r in self.results if not r['has_info']]

    def get_files_with_info(self) -> List[Dict]:
        """获取成功提取信息的文件"""
        return [r for r in self.results if r['has_info']]


def extract_from_docx(docx_path: str) -> Dict[str, Optional[str]]:
    """从docx文件提取信息的便捷函数"""
    reader = DocxReader(docx_path)
    return reader.extract_all_info()


if __name__ == "__main__":
    print("=== Docx读取模块测试 ===")
    print("用法: reader = DocxReader('实验报告.docx')")
    print("      info = reader.extract_all_info()")
    print("      print(info)")
